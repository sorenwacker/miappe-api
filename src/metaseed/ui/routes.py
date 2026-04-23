"""HTMX route handlers for the UI.

Provides FastAPI routes with Jinja2 templates for the HTMX-based interface.
"""

from __future__ import annotations

import contextlib
from io import BytesIO
from pathlib import Path
from typing import Any

from fastapi import FastAPI, File, HTTPException, Query, Request, UploadFile
from fastapi.responses import HTMLResponse, JSONResponse, RedirectResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from pydantic import ValidationError

from metaseed.facade import ProfileFacade
from metaseed.profiles import ProfileFactory
from metaseed.validators import validate as validate_data

from .helpers import (
    build_breadcrumb,
    build_inline_tables,
    collect_entities_by_type,
    collect_form_values,
    error_response,
    format_table_rows,
    format_validation_errors,
    get_field_data,
    get_items_store,
    get_parent_id_fields,
    get_reference_fields,
    get_table_column_info,
    get_table_columns,
    is_nested_field,
)
from .state import AppState, NestedEditContext

UI_DIR = Path(__file__).parent
TEMPLATES_DIR = UI_DIR / "templates"
STATIC_DIR = UI_DIR / "static"


def _get_profile_display_info(factory: ProfileFactory) -> list[dict]:
    """Get display information for all available profiles.

    Reads metadata from profile.yaml files.

    Args:
        factory: ProfileFactory instance.

    Returns:
        List of profile info dicts with name, display_name, description, root_entity, and versions.
    """
    from metaseed.specs.loader import SpecLoader

    profiles = []
    for name in factory.list_profiles():
        loader = SpecLoader(profile=name)
        versions = loader.list_versions(name)
        if not versions:
            continue

        # Load latest version to get profile metadata
        latest_version = versions[-1]
        try:
            profile_spec = loader.load_profile(latest_version, name)
            profiles.append(
                {
                    "name": name,
                    "display_name": profile_spec.display_name or name.upper(),
                    "description": profile_spec.description or f"{name} metadata profile.",
                    "root_entity": profile_spec.root_entity,
                    "versions": versions,
                    "latest_version": latest_version,
                }
            )
        except Exception:
            # Fallback if profile can't be loaded
            profiles.append(
                {
                    "name": name,
                    "display_name": name.upper(),
                    "description": f"{name} metadata profile.",
                    "root_entity": "Investigation",
                    "versions": versions,
                    "latest_version": latest_version,
                }
            )
    return profiles


def create_app(state: AppState | None = None) -> FastAPI:
    """Create the FastAPI application with HTMX routes.

    Args:
        state: Optional initial state. Creates new state if not provided.

    Returns:
        Configured FastAPI application.
    """
    app = FastAPI(title="Metaseed")

    if state is None:
        state = AppState()

    app.state.ui_state = state

    templates = Jinja2Templates(directory=str(TEMPLATES_DIR))

    # Add custom filter for formatting display values
    def format_display(value: Any) -> str:
        """Format a value for display in table cells."""
        if value is None:
            return ""
        if isinstance(value, list):
            # Filter out empty values and join
            return ", ".join(str(v) for v in value if v)
        # Handle string representation of empty list
        if isinstance(value, str) and value.strip() in ("[]", ""):
            return ""
        return str(value)

    templates.env.filters["display"] = format_display

    app.mount("/static", StaticFiles(directory=str(STATIC_DIR)), name="static")

    def get_state() -> AppState:
        return app.state.ui_state

    # Mount spec builder routes
    from .spec_builder_routes import create_spec_builder_router

    spec_builder_router = create_spec_builder_router(templates, get_state)
    app.include_router(spec_builder_router)

    # -------------------------------------------------------------------------
    # Main routes
    # -------------------------------------------------------------------------

    @app.get("/", response_class=HTMLResponse)
    async def index(request: Request) -> HTMLResponse:
        """Render the main page."""
        state = get_state()
        facade = state.get_or_create_facade()
        profile_factory = ProfileFactory()

        # Get editing node info if set
        editing_node = None
        if state.editing_node_id:
            editing_node = state.nodes_by_id.get(state.editing_node_id)

        return templates.TemplateResponse(
            request,
            "base.html",
            {
                "profiles": profile_factory.list_profiles(),
                "current_profile": state.profile,
                "version": facade.version,
                "root_types": state.get_root_entity_types()[:3],
                "tree_nodes": state.get_tree_data(),
                "editing_node_id": state.editing_node_id,
                "editing_node_type": editing_node.entity_type if editing_node else None,
            },
        )

    @app.get("/profile/{name}")
    async def switch_profile(name: str) -> RedirectResponse:
        """Switch to a different profile."""
        state = get_state()
        profile_factory = ProfileFactory()

        if name not in profile_factory.list_profiles():
            raise HTTPException(status_code=400, detail=f"Unknown profile: {name}")

        state.profile = name
        state.facade = None
        state.reset()

        return RedirectResponse(url="/", status_code=303)

    @app.get("/load-example/{profile_name}/{version}")
    async def load_example(profile_name: str, version: str) -> RedirectResponse:
        """Load example data for a profile version."""
        import yaml

        from metaseed.models import get_model

        state = get_state()
        profile_factory = ProfileFactory()

        if profile_name not in profile_factory.list_profiles():
            raise HTTPException(status_code=400, detail=f"Unknown profile: {profile_name}")

        # Find example file (structure: examples/profile/version/*.yaml)
        examples_dir = Path(__file__).parent.parent.parent.parent / "examples"
        version_dir = examples_dir / profile_name / version

        if not version_dir.exists():
            raise HTTPException(
                status_code=404, detail=f"No example available for {profile_name} v{version}"
            )

        yaml_files = list(version_dir.glob("*.yaml"))
        if not yaml_files:
            raise HTTPException(status_code=404, detail=f"No example file found in {version_dir}")

        example_file = yaml_files[0]

        # Load example data
        try:
            example_data = yaml.safe_load(example_file.read_text(encoding="utf-8"))
        except yaml.YAMLError as e:
            raise HTTPException(status_code=500, detail=f"Error loading example: {e}") from e

        # Reset state and set profile
        state.reset()
        state.profile = profile_name
        state.version = version
        state.facade = None
        facade = state.get_or_create_facade()

        # Determine root entity type from spec
        from metaseed.specs.loader import SpecLoader

        loader = SpecLoader(profile=profile_name)
        spec = loader.load_profile(version, profile_name)
        root_entity = spec.root_entity or "Investigation"

        # Create entity instance from example data
        try:
            Model = get_model(root_entity, version, profile=profile_name)
            instance = Model(**example_data)
        except Exception as e:
            raise HTTPException(
                status_code=500, detail=f"Error creating entity from example: {e}"
            ) from e

        # Add to tree and set as editing
        node = state.add_node(root_entity, instance)
        state.editing_node_id = node.id

        # Populate current_nested_items from the instance
        helper = getattr(facade, root_entity)
        for field_name in helper.nested_fields:
            if example_data.get(field_name):
                items = example_data[field_name]
                if isinstance(items, list):
                    state.current_nested_items[field_name] = list(items)
                elif isinstance(items, dict):
                    state.current_nested_items[field_name] = [items]

        # Redirect to main page which will show the loaded entity
        return RedirectResponse(url="/", status_code=303)

    @app.post("/reset", response_class=HTMLResponse)
    async def reset_state() -> HTMLResponse:
        """Reset all application state. Used for testing."""
        state = get_state()
        state.reset()
        return HTMLResponse(content="OK")

    # -------------------------------------------------------------------------
    # Entity form routes
    # -------------------------------------------------------------------------

    @app.get("/form/{entity_type}", response_class=HTMLResponse)
    async def new_entity_form(request: Request, entity_type: str) -> HTMLResponse:
        """Render a new entity form."""
        state = get_state()
        profile_factory = ProfileFactory()

        # Check if profile is specified in query params
        profile = request.query_params.get("profile")

        # Show profile selection if this is a root entity request without profile
        if not profile:
            profiles_info = _get_profile_display_info(profile_factory)
            root_entities = {p["root_entity"] for p in profiles_info}
            if entity_type in root_entities:
                return templates.TemplateResponse(
                    request,
                    "partials/profile_select.html",
                    {"profiles": profiles_info},
                )

        # If profile specified, switch to it
        version = request.query_params.get("version")
        if profile and profile in profile_factory.list_profiles():
            state.profile = profile
            state.version = version  # None means use latest
            state.facade = None  # Reset facade to use new profile

        facade = state.get_or_create_facade()

        try:
            helper = getattr(facade, entity_type)
        except AttributeError as e:
            raise HTTPException(
                status_code=404, detail=f"Entity type not found: {entity_type}"
            ) from e

        state.editing_node_id = None
        state.current_nested_items = {}  # Reset nested items for new entity

        fields = get_field_data(helper)

        # Auto-populate values for certain fields
        auto_values = {}
        if "miappe_version" in helper.all_fields:
            auto_values["miappe_version"] = facade.version

        # Check if example exists for this profile/version
        examples_dir = Path(__file__).parent.parent.parent.parent / "examples"
        example_exists = (examples_dir / state.profile / facade.version).exists()

        return templates.TemplateResponse(
            request,
            "partials/form.html",
            {
                "entity_type": entity_type,
                "is_edit": False,
                "node_id": None,
                "description": helper.description,
                "ontology_term": helper.ontology_term,
                "required_fields": [f for f in fields if f["required"]],
                "optional_fields": [
                    f for f in fields if not f["required"] and not is_nested_field(f)
                ],
                "nested_fields": [f for f in fields if is_nested_field(f)],
                "values": auto_values,
                "auto_fields": set(auto_values.keys()),
                "current_profile": state.profile,
                "current_version": facade.version,
                "example_available": example_exists,
            },
        )

    @app.get("/form/{entity_type}/{node_id}", response_class=HTMLResponse)
    async def edit_entity_form(request: Request, entity_type: str, node_id: str) -> HTMLResponse:
        """Render an edit form for an existing entity."""
        state = get_state()
        facade = state.get_or_create_facade()

        node = state.nodes_by_id.get(node_id)
        if not node:
            raise HTTPException(status_code=404, detail=f"Node not found: {node_id}")

        try:
            helper = getattr(facade, entity_type)
        except AttributeError as e:
            raise HTTPException(
                status_code=404, detail=f"Entity type not found: {entity_type}"
            ) from e

        state.editing_node_id = node_id
        state.nested_edit_stack = []  # Clear nested context when editing root entity

        fields = get_field_data(helper)
        values = {}
        if node.instance and hasattr(node.instance, "model_dump"):
            values = node.instance.model_dump(exclude_none=True)

        # Merge current_nested_items (from table editing) into values for display
        for field_name, items in state.current_nested_items.items():
            if items:
                values[field_name] = items

        # Only load nested items from entity if we don't have any pending edits
        if not state.current_nested_items:
            for field_name in helper.nested_fields:
                if values.get(field_name):
                    items = values[field_name]
                    if isinstance(items, list):
                        state.current_nested_items[field_name] = [
                            item.model_dump() if hasattr(item, "model_dump") else item
                            for item in items
                        ]

        # Auto-populate values for certain fields
        auto_fields = set()
        if "miappe_version" in helper.all_fields:
            values["miappe_version"] = facade.version
            auto_fields.add("miappe_version")

        # Build inline table data for nested fields
        inline_tables = build_inline_tables(state, facade, entity_type)

        return templates.TemplateResponse(
            request,
            "partials/form.html",
            {
                "entity_type": entity_type,
                "is_edit": True,
                "node_id": node_id,
                "description": helper.description,
                "ontology_term": helper.ontology_term,
                "required_fields": [f for f in fields if f["required"]],
                "optional_fields": [
                    f for f in fields if not f["required"] and not is_nested_field(f)
                ],
                "nested_fields": [f for f in fields if is_nested_field(f)],
                "values": values,
                "auto_fields": auto_fields,
                "inline_tables": inline_tables,
            },
        )

    # -------------------------------------------------------------------------
    # Entity CRUD routes
    # -------------------------------------------------------------------------

    @app.post("/entity", response_class=HTMLResponse)
    async def create_entity(request: Request) -> HTMLResponse:
        """Create a new entity."""
        state = get_state()
        facade = state.get_or_create_facade()

        form_data = await request.form()
        entity_type = form_data.get("_entity_type")

        if not entity_type:
            return error_response(request, templates, "Entity type is required")

        try:
            helper = getattr(facade, entity_type)
        except AttributeError:
            return error_response(request, templates, f"Unknown entity type: {entity_type}")

        values = collect_form_values(dict(form_data), helper)

        try:
            instance = helper.create(**values)
            node = state.add_node(entity_type, instance)
            state.editing_node_id = node.id

            # Load nested items from the new entity for editing
            state.current_nested_items = {}
            if hasattr(instance, "model_dump"):
                data = instance.model_dump(exclude_none=True)
                for field_name in helper.nested_fields:
                    if data.get(field_name):
                        items = data[field_name]
                        if isinstance(items, list):
                            state.current_nested_items[field_name] = [
                                item.model_dump() if hasattr(item, "model_dump") else item
                                for item in items
                            ]

            # Return the edit form for the newly created entity
            return _render_entity_form(
                request,
                templates,
                facade,
                helper,
                entity_type,
                node.id,
                instance,
                f"Created {entity_type}: {node.label}",
                state,
            )

        except ValidationError as e:
            return _render_form_with_errors(
                request, templates, facade, helper, entity_type, None, values, e
            )

    @app.put("/entity/{node_id}", response_class=HTMLResponse)
    async def update_entity(request: Request, node_id: str) -> HTMLResponse:
        """Update an existing entity."""
        state = get_state()
        facade = state.get_or_create_facade()

        node = state.nodes_by_id.get(node_id)
        if not node:
            return error_response(request, templates, f"Node not found: {node_id}")

        form_data = await request.form()
        entity_type = node.entity_type

        try:
            helper = getattr(facade, entity_type)
        except AttributeError:
            return error_response(request, templates, f"Unknown entity type: {entity_type}")

        values = collect_form_values(dict(form_data), helper)

        # Merge nested items from table editing
        for field_name, items in state.current_nested_items.items():
            if field_name in helper.nested_fields and items:
                cleaned_items = []
                for item in items:
                    if isinstance(item, dict):
                        cleaned = {k: v for k, v in item.items() if not k.startswith("_") and v}
                        if cleaned:
                            cleaned_items.append(cleaned)
                if cleaned_items:
                    values[field_name] = cleaned_items

        try:
            instance = helper.create(**values)
            state.update_node(node_id, instance)

            # Reload nested items from saved instance
            state.current_nested_items = {}
            if hasattr(instance, "model_dump"):
                data = instance.model_dump(exclude_none=True)
                for field_name in helper.nested_fields:
                    if data.get(field_name):
                        items = data[field_name]
                        if isinstance(items, list):
                            state.current_nested_items[field_name] = [
                                item.model_dump() if hasattr(item, "model_dump") else item
                                for item in items
                            ]

            return _render_entity_form(
                request,
                templates,
                facade,
                helper,
                entity_type,
                node_id,
                instance,
                f"Saved {entity_type}: {node.label}",
                state,
            )

        except ValidationError as e:
            return _render_form_with_errors(
                request, templates, facade, helper, entity_type, node_id, values, e
            )

    @app.delete("/entity/{node_id}", response_class=HTMLResponse)
    async def delete_entity(request: Request, node_id: str) -> HTMLResponse:
        """Delete an entity."""
        state = get_state()

        node = state.nodes_by_id.get(node_id)
        if not node:
            return error_response(request, templates, f"Node not found: {node_id}")

        entity_type = node.entity_type
        label = node.label

        state.delete_node(node_id)

        return templates.TemplateResponse(
            request,
            "index.html",
            {
                "notification": {
                    "type": "warning",
                    "message": f"Deleted {entity_type}: {label}",
                },
            },
        )

    # -------------------------------------------------------------------------
    # Table routes (nested entity lists)
    # -------------------------------------------------------------------------

    @app.get("/table/{entity_type}/{field_name}", response_class=HTMLResponse)
    async def table_view(request: Request, entity_type: str, field_name: str) -> HTMLResponse:
        """Render the nested table view for a list field."""
        state = get_state()
        facade = state.get_or_create_facade()

        try:
            helper = getattr(facade, entity_type)
        except AttributeError as e:
            raise HTTPException(
                status_code=404, detail=f"Entity type not found: {entity_type}"
            ) from e

        # Clear nested stack if viewing root entity's table
        if state.editing_node_id:
            root_node = state.nodes_by_id.get(state.editing_node_id)
            if root_node and root_node.entity_type == entity_type:
                state.nested_edit_stack = []

        nested_fields = helper.nested_fields
        if field_name not in nested_fields:
            raise HTTPException(status_code=404, detail=f"Field not found: {field_name}")

        nested_entity_type = nested_fields[field_name]
        col_info = get_table_column_info(facade, nested_entity_type)
        _, items = get_items_store(state, entity_type, field_name)

        # Get reference fields for the nested entity type
        reference_fields = get_reference_fields(
            profile=state.profile,
            version=facade.version,
            entity_type=nested_entity_type,
        )

        # Get parent ID fields (hidden from display, used for serialization only)
        parent_id_fields = get_parent_id_fields(reference_fields, entity_type)
        display_columns = [c for c in col_info["columns"] if c not in parent_id_fields]

        rows = []
        for i, item in enumerate(items):
            if hasattr(item, "model_dump"):
                row = item.model_dump(exclude_none=True)
            elif isinstance(item, dict):
                row = item.copy()
            else:
                row = {"value": str(item)}
            row["_idx"] = i
            rows.append(row)

        return templates.TemplateResponse(
            request,
            "partials/table.html",
            {
                "field_name": field_name,
                "entity_type": nested_entity_type,
                "columns": display_columns,
                "column_types": col_info["column_types"],
                "column_constraints": col_info["column_constraints"],
                "required_columns": col_info["required_columns"],
                "has_nested_children": col_info["has_nested_children"],
                "reference_fields": reference_fields,
                "rows": rows,
                "parent_entity_type": entity_type,
                "editing_node_id": state.editing_node_id,
                "breadcrumb": build_breadcrumb(state),
                "nested_context": state.nested_edit_stack[-1] if state.nested_edit_stack else None,
            },
        )

    @app.post("/table/{parent_entity_type}/{field_name}/row", response_class=HTMLResponse)
    async def add_table_row(
        request: Request, parent_entity_type: str, field_name: str
    ) -> HTMLResponse:
        """Add a new row to the nested table."""
        state = get_state()
        facade = state.get_or_create_facade()

        _, items = get_items_store(state, parent_entity_type, field_name)

        parent_helper = getattr(facade, parent_entity_type, None)
        entity_type = parent_helper.nested_fields.get(field_name) if parent_helper else None
        col_info = get_table_column_info(facade, entity_type)

        # Get reference fields for this entity type
        reference_fields = (
            get_reference_fields(
                profile=state.profile,
                version=facade.version,
                entity_type=entity_type,
            )
            if entity_type
            else {}
        )

        # Find fields that reference the parent and should be auto-filled
        parent_id_fields = get_parent_id_fields(reference_fields, parent_entity_type)

        # Get parent's identifier to auto-fill parent reference fields
        parent_identifier = ""
        if parent_id_fields and state.editing_node_id:
            node = state.nodes_by_id.get(state.editing_node_id)
            if node:
                # Get the actual parent data - could be the root node or from nested stack
                parent_data = None
                if node.entity_type == parent_entity_type:
                    # Parent is the root node
                    if hasattr(node.instance, "model_dump"):
                        parent_data = node.instance.model_dump(exclude_none=True)
                elif state.nested_edit_stack:
                    # Check nested stack for parent
                    for ctx in reversed(state.nested_edit_stack):
                        if ctx.entity_type == parent_entity_type:
                            # Get the item from nested items
                            parent_items = state.current_nested_items.get(ctx.field_name, [])
                            if ctx.row_idx < len(parent_items):
                                parent_data = parent_items[ctx.row_idx]
                            break

                if parent_data:
                    # Get the identifier field value (usually "unique_id" or "identifier")
                    for target_field in parent_id_fields.values():
                        if target_field in parent_data:
                            parent_identifier = str(parent_data[target_field])
                            break

        new_row = dict.fromkeys(col_info["columns"], "")
        new_row["_idx"] = len(items)

        # Auto-fill parent reference fields
        for field_name_ref in parent_id_fields:
            if field_name_ref in new_row:
                new_row[field_name_ref] = parent_identifier

        items.append(new_row)

        # Check if this is from an inline table (hx-target contains 'inline')
        hx_target = request.headers.get("hx-target", "")
        is_inline = "inline" in hx_target

        template_name = "partials/inline_table_row.html" if is_inline else "partials/table_row.html"

        # Filter out parent ID fields from display (data kept for serialization)
        display_columns = [c for c in col_info["columns"] if c not in parent_id_fields]

        return templates.TemplateResponse(
            request,
            template_name,
            {
                "row": new_row,
                "columns": display_columns,
                "column_types": col_info["column_types"],
                "column_constraints": col_info["column_constraints"],
                "reference_fields": reference_fields,
                "parent_id_fields": parent_id_fields,
                "field_name": field_name,
                "parent_entity_type": parent_entity_type,
                "entity_type": entity_type,
                "has_nested_children": col_info["has_nested_children"],
            },
        )

    @app.delete("/table/{parent_entity_type}/{field_name}/row/{idx}", response_class=HTMLResponse)
    async def delete_table_row(parent_entity_type: str, field_name: str, idx: int) -> HTMLResponse:
        """Delete a row from the nested table."""
        state = get_state()
        _, items = get_items_store(state, parent_entity_type, field_name)

        if 0 <= idx < len(items):
            del items[idx]
            for i, item in enumerate(items):
                if isinstance(item, dict):
                    item["_idx"] = i

        return HTMLResponse(content="")

    @app.post(
        "/table/{parent_entity_type}/{field_name}/row/{idx}/cell", response_class=HTMLResponse
    )
    async def update_table_cell(
        request: Request, parent_entity_type: str, field_name: str, idx: int
    ) -> HTMLResponse:
        """Update a cell value in the nested table."""
        state = get_state()
        _, items = get_items_store(state, parent_entity_type, field_name)

        if 0 <= idx < len(items):
            form_data = await request.form()
            item = items[idx]
            if isinstance(item, dict):
                for key, value in form_data.items():
                    if not key.startswith("_"):
                        item[key] = value

        return HTMLResponse(content="")

    @app.post("/table/{parent_entity_type}/{field_name}/bulk", response_class=HTMLResponse)
    async def bulk_update_rows(
        request: Request, parent_entity_type: str, field_name: str
    ) -> HTMLResponse:
        """Bulk update multiple rows with the same value."""
        state = get_state()
        facade = state.get_or_create_facade()

        form_data = await request.form()
        field = form_data.get("bulk-edit-field", "")
        value = form_data.get("bulk-edit-value", "")
        indices_str = form_data.get("indices", "")

        if not field or not indices_str:
            return error_response(request, templates, "Field and indices are required")

        try:
            indices = [int(i.strip()) for i in indices_str.split(",") if i.strip()]
        except ValueError:
            return error_response(request, templates, "Invalid indices format")

        _, items = get_items_store(state, parent_entity_type, field_name)

        updated_count = 0
        for idx in indices:
            if 0 <= idx < len(items):
                item = items[idx]
                if isinstance(item, dict):
                    item[field] = value
                    updated_count += 1

        # Return the refreshed table view
        try:
            helper = getattr(facade, parent_entity_type)
        except AttributeError as e:
            raise HTTPException(
                status_code=404, detail=f"Entity type not found: {parent_entity_type}"
            ) from e

        nested_entity_type = helper.nested_fields.get(field_name)
        col_info = get_table_column_info(facade, nested_entity_type)

        # Get reference fields and filter out parent ID fields from display
        reference_fields = get_reference_fields(
            profile=state.profile,
            version=facade.version,
            entity_type=nested_entity_type,
        )
        parent_id_fields = get_parent_id_fields(reference_fields, parent_entity_type)
        display_columns = [c for c in col_info["columns"] if c not in parent_id_fields]

        rows = []
        for i, item in enumerate(items):
            if hasattr(item, "model_dump"):
                row = item.model_dump(exclude_none=True)
            elif isinstance(item, dict):
                row = item.copy()
            else:
                row = {"value": str(item)}
            row["_idx"] = i
            rows.append(row)

        response = templates.TemplateResponse(
            request,
            "partials/table.html",
            {
                "field_name": field_name,
                "entity_type": nested_entity_type,
                "columns": display_columns,
                "column_types": col_info["column_types"],
                "column_constraints": col_info["column_constraints"],
                "required_columns": col_info["required_columns"],
                "has_nested_children": col_info["has_nested_children"],
                "rows": rows,
                "parent_entity_type": parent_entity_type,
                "editing_node_id": state.editing_node_id,
                "breadcrumb": build_breadcrumb(state),
                "nested_context": state.nested_edit_stack[-1] if state.nested_edit_stack else None,
                "notification": {
                    "type": "success",
                    "message": f"Updated {updated_count} rows",
                },
            },
        )
        return response

    @app.post("/table/{parent_entity_type}/{field_name}/paste", response_class=HTMLResponse)
    async def paste_cells(
        request: Request, parent_entity_type: str, field_name: str
    ) -> HTMLResponse:
        """Apply pasted cell values from clipboard."""
        import json

        state = get_state()
        form_data = await request.form()
        changes_json = form_data.get("changes", "[]")

        try:
            changes = json.loads(changes_json)
        except json.JSONDecodeError:
            return error_response(request, templates, "Invalid paste data format")

        _, items = get_items_store(state, parent_entity_type, field_name)

        updated_count = 0
        for change in changes:
            idx = change.get("idx")
            field = change.get("field")
            value = change.get("value")

            if idx is not None and field and 0 <= idx < len(items):
                item = items[idx]
                if isinstance(item, dict):
                    item[field] = value
                    updated_count += 1

        return templates.TemplateResponse(
            request,
            "components/notification.html",
            {
                "type": "success",
                "message": f"Pasted {updated_count} cells",
            },
        )

    # -------------------------------------------------------------------------
    # Lookup API routes (for cross-entity references)
    # -------------------------------------------------------------------------

    @app.get("/api/lookup/{entity_type}")
    async def lookup_entities(
        entity_type: str,
        q: str = Query(default="", description="Search query"),
    ) -> JSONResponse:
        """Search entities of a given type for autocomplete.

        Args:
            entity_type: The type of entity to search (e.g., "ObservationUnit").
            q: Search query to filter by identifier and label fields.

        Returns:
            JSON with results list containing value and label for each match.
        """
        state = get_state()
        facade = state.get_or_create_facade()

        # Collect all entities organized by type
        entities_by_type = collect_entities_by_type(state, facade)

        # Get entities of the requested type
        entities = entities_by_type.get(entity_type, [])

        # Filter by search query
        query = q.lower().strip()
        if query:
            filtered = []
            for entity in entities:
                value = entity.get("value", "").lower()
                label = entity.get("label", "").lower()
                if query in value or query in label:
                    filtered.append(entity)
            entities = filtered

        # Return unique results (dedupe by value)
        seen = set()
        results = []
        for entity in entities:
            value = entity.get("value", "")
            if value and value not in seen:
                seen.add(value)
                results.append(
                    {
                        "value": value,
                        "label": entity.get("label", value),
                    }
                )

        return JSONResponse(content={"results": results})

    @app.get("/api/reference-fields/{entity_type}")
    async def get_reference_fields_api(entity_type: str) -> JSONResponse:
        """Get reference field definitions for an entity type.

        Args:
            entity_type: The entity type to get reference fields for.

        Returns:
            JSON with reference fields mapping field name to target info.
        """
        state = get_state()
        facade = state.get_or_create_facade()

        ref_fields = get_reference_fields(
            profile=state.profile,
            version=facade.version,
            entity_type=entity_type,
        )

        return JSONResponse(content=ref_fields)

    @app.get("/api/graph")
    async def get_graph() -> JSONResponse:
        """Return graph data for visualization.

        Builds nodes and edges from the entity tree for vis.js network graph.

        Returns:
            JSON with 'nodes' and 'edges' lists.
        """
        from metaseed.ui.services.graph import build_graph

        state = get_state()
        return JSONResponse(content=build_graph(state))

    # -------------------------------------------------------------------------
    # Nested form routes
    # -------------------------------------------------------------------------

    @app.get("/nested/{parent_type}/{field_name}/{idx}", response_class=HTMLResponse)
    async def edit_nested_item(
        request: Request, parent_type: str, field_name: str, idx: int
    ) -> HTMLResponse:
        """Edit a nested item (e.g., a Study within an Investigation)."""
        state = get_state()
        facade = state.get_or_create_facade()
        is_resume = request.query_params.get("resume") == "true"

        try:
            parent_helper = getattr(facade, parent_type)
        except AttributeError as e:
            raise HTTPException(
                status_code=404, detail=f"Entity type not found: {parent_type}"
            ) from e

        if field_name not in parent_helper.nested_fields:
            raise HTTPException(status_code=404, detail=f"Field not found: {field_name}")

        nested_entity_type = parent_helper.nested_fields[field_name]
        nested_helper = getattr(facade, nested_entity_type, None)
        _, items = get_items_store(state, parent_type, field_name)

        if idx < 0 or idx >= len(items):
            raise HTTPException(status_code=404, detail=f"Row not found: {idx}")

        # Handle resume vs fresh edit
        if is_resume and state.nested_edit_stack:
            context = state.nested_edit_stack[-1]
            item_data = items[idx]
            if hasattr(item_data, "model_dump"):
                item_data = item_data.model_dump(exclude_none=True)
            elif isinstance(item_data, dict):
                item_data = item_data.copy()
            else:
                item_data = {}
            for nf, nv in context.nested_items.items():
                item_data[nf] = nv
        else:
            item_data = items[idx]
            if hasattr(item_data, "model_dump"):
                item_data = item_data.model_dump(exclude_none=True)
            elif isinstance(item_data, dict):
                item_data = item_data.copy()
            else:
                item_data = {}

            context = NestedEditContext(
                field_name=field_name,
                row_idx=idx,
                entity_type=nested_entity_type,
                parent_entity_type=parent_type,
            )

            if nested_helper:
                for nested_field in nested_helper.nested_fields:
                    if item_data.get(nested_field):
                        nested_items = item_data[nested_field]
                        if isinstance(nested_items, list):
                            context.nested_items[nested_field] = [
                                i.model_dump() if hasattr(i, "model_dump") else i
                                for i in nested_items
                            ]

            state.nested_edit_stack.append(context)

        fields = get_field_data(nested_helper) if nested_helper else []
        values = item_data.copy() if isinstance(item_data, dict) else {}
        if state.nested_edit_stack:
            ctx = state.nested_edit_stack[-1]
            for nf, nv in ctx.nested_items.items():
                values[nf] = nv

        # Build inline tables for this nested entity's own nested fields
        inline_tables = {}
        if state.nested_edit_stack:
            ctx = state.nested_edit_stack[-1]
            inline_tables = build_inline_tables(
                state, facade, nested_entity_type, items_source=ctx.nested_items
            )

        return templates.TemplateResponse(
            request,
            "partials/nested_form.html",
            {
                "entity_type": nested_entity_type,
                "parent_entity_type": parent_type,
                "field_name": field_name,
                "row_idx": idx,
                "description": nested_helper.description if nested_helper else "",
                "ontology_term": nested_helper.ontology_term if nested_helper else "",
                "required_fields": [f for f in fields if f["required"]],
                "optional_fields": [
                    f for f in fields if not f["required"] and not is_nested_field(f)
                ],
                "nested_fields": [f for f in fields if is_nested_field(f)],
                "values": values,
                "auto_fields": set(),
                "editing_node_id": state.editing_node_id,
                "breadcrumb": build_breadcrumb(state),
                "inline_tables": inline_tables,
            },
        )

    @app.post("/nested/{parent_type}/{field_name}/{idx}", response_class=HTMLResponse)
    async def save_nested_item(
        request: Request, parent_type: str, field_name: str, idx: int
    ) -> HTMLResponse:
        """Save changes to a nested item."""
        state = get_state()
        facade = state.get_or_create_facade()

        _, items = get_items_store(state, parent_type, field_name)
        if idx < 0 or idx >= len(items):
            raise HTTPException(status_code=404, detail=f"Row not found: {idx}")

        form_data = await request.form()
        go_back = form_data.get("_action") == "back"

        parent_helper = getattr(facade, parent_type, None)
        nested_entity_type = parent_helper.nested_fields.get(field_name) if parent_helper else None
        nested_helper = getattr(facade, nested_entity_type, None) if nested_entity_type else None

        item = items[idx]
        if isinstance(item, dict):
            for key, value in form_data.items():
                if not key.startswith("_"):
                    if nested_helper:
                        info = (
                            nested_helper.field_info(key) if key in nested_helper.all_fields else {}
                        )
                        if info.get("type") == "integer" and value:
                            with contextlib.suppress(ValueError):
                                value = int(value)
                        elif info.get("type") == "float" and value:
                            with contextlib.suppress(ValueError):
                                value = float(value)
                    if value:
                        item[key] = value

            if state.nested_edit_stack:
                context = state.nested_edit_stack[-1]
                for nested_field, nested_values in context.nested_items.items():
                    if nested_values:
                        item[nested_field] = nested_values

        if go_back:
            if state.nested_edit_stack:
                state.nested_edit_stack.pop()

            # Filter out parent ID fields from display
            all_columns = get_table_columns(facade, nested_entity_type)
            reference_fields = get_reference_fields(
                profile=state.profile,
                version=facade.version,
                entity_type=nested_entity_type,
            )
            parent_id_fields = get_parent_id_fields(reference_fields, parent_type)
            display_columns = [c for c in all_columns if c not in parent_id_fields]

            return templates.TemplateResponse(
                request,
                "partials/table.html",
                {
                    "field_name": field_name,
                    "entity_type": nested_entity_type,
                    "columns": display_columns,
                    "rows": format_table_rows(items),
                    "parent_entity_type": parent_type,
                    "editing_node_id": state.editing_node_id,
                },
            )

        fields = get_field_data(nested_helper) if nested_helper else []
        values = item.copy() if isinstance(item, dict) else {}

        if state.nested_edit_stack:
            ctx = state.nested_edit_stack[-1]
            for nf, nv in ctx.nested_items.items():
                values[nf] = nv

        # Build inline tables for this nested entity's own nested fields
        inline_tables = {}
        if state.nested_edit_stack:
            ctx = state.nested_edit_stack[-1]
            inline_tables = build_inline_tables(
                state, facade, nested_entity_type, items_source=ctx.nested_items
            )

        return templates.TemplateResponse(
            request,
            "partials/nested_form.html",
            {
                "entity_type": nested_entity_type,
                "parent_entity_type": parent_type,
                "field_name": field_name,
                "row_idx": idx,
                "description": nested_helper.description if nested_helper else "",
                "ontology_term": nested_helper.ontology_term if nested_helper else "",
                "required_fields": [f for f in fields if f["required"]],
                "optional_fields": [
                    f for f in fields if not f["required"] and not is_nested_field(f)
                ],
                "nested_fields": [f for f in fields if is_nested_field(f)],
                "values": values,
                "auto_fields": set(),
                "editing_node_id": state.editing_node_id,
                "breadcrumb": build_breadcrumb(state),
                "inline_tables": inline_tables,
                "success_message": "Saved",
            },
        )

    # -------------------------------------------------------------------------
    # Export and validation routes
    # -------------------------------------------------------------------------

    @app.get("/export")
    async def export_excel(_request: Request) -> StreamingResponse:
        """Export current entity data to Excel file."""
        from openpyxl import Workbook

        state = get_state()
        facade = state.get_or_create_facade()

        wb = Workbook()
        wb.remove(wb.active)

        entities_by_type: dict[str, list[dict]] = {}

        def extract_nested_entities(data: dict, entity_type: str) -> None:
            helper = getattr(facade, entity_type, None)
            if not helper:
                return

            for field_name, nested_type in helper.nested_fields.items():
                if data.get(field_name):
                    nested_items = data[field_name]
                    if isinstance(nested_items, list):
                        for item in nested_items:
                            if hasattr(item, "model_dump"):
                                item_data = item.model_dump(exclude_none=True)
                            elif isinstance(item, dict):
                                item_data = item.copy()
                            else:
                                continue

                            if nested_type not in entities_by_type:
                                entities_by_type[nested_type] = []
                            entities_by_type[nested_type].append(item_data)
                            extract_nested_entities(item_data, nested_type)

        for node in state.nodes_by_id.values():
            entity_type = node.entity_type
            if entity_type not in entities_by_type:
                entities_by_type[entity_type] = []

            if hasattr(node.instance, "model_dump"):
                data = node.instance.model_dump(exclude_none=True)
            else:
                data = {}

            entities_by_type[entity_type].append(data)
            extract_nested_entities(data, entity_type)

        # Create sheets for ALL entity types in the profile, even if empty
        for entity_type in facade.entities:
            helper = getattr(facade, entity_type, None)
            if not helper:
                continue

            ws = wb.create_sheet(entity_type)
            nested_fields = set(helper.nested_fields.keys())
            # Include all fields - nested fields will show counts
            columns = helper.all_fields

            ws.append(columns)

            # Add data rows if we have any
            entities = entities_by_type.get(entity_type, [])
            for entity_data in entities:
                row = []
                for col in columns:
                    value = entity_data.get(col, "")
                    if col in nested_fields:
                        # Show count for nested fields
                        if isinstance(value, list):
                            value = len(value)
                        elif value:
                            value = 1
                        else:
                            value = 0
                    elif isinstance(value, list):
                        if value and not isinstance(value[0], dict):
                            value = ", ".join(str(v) for v in value)
                        else:
                            value = len(value)
                    elif isinstance(value, dict):
                        value = "[object]"
                    elif not isinstance(value, str | int | float | bool | type(None)):
                        # Convert Pydantic types like AnyUrl to string
                        value = str(value)
                    row.append(value)
                ws.append(row)

        output = BytesIO()
        wb.save(output)
        output.seek(0)

        # Generate filename: YYMMDD-<standard>-<version>-<id>.xlsx
        from datetime import datetime

        date_str = datetime.now().strftime("%y%m%d")
        version_str = facade.version.replace(".", "-")

        # Try to get ID from root entity
        entity_id = "export"
        root_nodes = [n for n in state.nodes_by_id.values() if n.parent_id is None]
        if root_nodes:
            root_node = root_nodes[0]
            if hasattr(root_node.instance, "model_dump"):
                root_data = root_node.instance.model_dump()
                if root_data.get("unique_id"):
                    # Sanitize the ID for filename
                    entity_id = str(root_data["unique_id"]).replace("/", "-").replace(":", "-")[:30]

        filename = f"{date_str}-{state.profile}-{version_str}-{entity_id}.xlsx"
        return StreamingResponse(
            output,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="{filename}"'},
        )

    @app.post("/import", response_class=HTMLResponse)
    async def import_isa(
        request: Request,
        file: UploadFile = File(...),
    ) -> HTMLResponse:
        """Import ISA-JSON file and create entities."""
        import tempfile

        from metaseed.importers.isa import ISAImporter

        state = get_state()
        facade = state.get_or_create_facade()

        # Check file type
        filename = file.filename or ""
        content = await file.read()

        try:
            importer = ISAImporter()

            if filename.endswith(".json"):
                # ISA-JSON file
                with tempfile.NamedTemporaryFile(mode="wb", suffix=".json", delete=False) as tmp:
                    tmp.write(content)
                    tmp_path = tmp.name

                result = importer.import_json(tmp_path)
                Path(tmp_path).unlink()
            else:
                return error_response(
                    request,
                    templates,
                    "Unsupported file type. Please upload an ISA-JSON file (.json).",
                )

            # Create Investigation entity
            if result.investigation and "Investigation" in facade.entities:
                helper = facade.Investigation
                inv_data = result.investigation.copy()

                # Add miappe_version for MIAPPE profile
                if state.profile == "miappe" and "miappe_version" in helper.all_fields:
                    inv_data["miappe_version"] = facade.version

                # Add nested items directly to the investigation data
                if result.studies and "studies" in helper.nested_fields:
                    inv_data["studies"] = result.studies

                if result.persons:
                    # Try different field names for persons
                    for field_name in ["persons", "contacts", "people"]:
                        if field_name in helper.nested_fields:
                            inv_data[field_name] = result.persons
                            break

                try:
                    instance = helper.create(**inv_data)
                    node = state.add_node("Investigation", instance)
                    state.editing_node_id = node.id
                    state.current_nested_items = {}

                    # Also store in current_nested_items for editing
                    if result.studies:
                        state.current_nested_items["studies"] = result.studies
                    if result.persons:
                        for field_name in ["persons", "contacts", "people"]:
                            if field_name in helper.nested_fields:
                                state.current_nested_items[field_name] = result.persons
                                break

                except ValidationError:
                    pass  # Skip invalid entities

            # Return success notification
            response = templates.TemplateResponse(
                request,
                "components/notification.html",
                {
                    "type": "success",
                    "message": result.summary,
                },
            )
            response.headers["HX-Trigger"] = "refreshPage"
            return response

        except Exception as e:
            return error_response(request, templates, f"Import failed: {e!s}")

    @app.post("/validate", response_class=HTMLResponse)
    async def validate_form(request: Request) -> HTMLResponse:
        """Validate form data against MIAPPE spec."""
        state = get_state()
        form_data = await request.form()
        entity_type = form_data.get("_entity_type")

        if not entity_type:
            return templates.TemplateResponse(
                request,
                "components/validation_result.html",
                {
                    "valid": False,
                    "errors": [
                        {
                            "field": "_entity_type",
                            "message": "Entity type is required",
                            "rule": "required",
                        }
                    ],
                },
            )

        facade = ProfileFacade(profile=state.profile)
        helper = getattr(facade, entity_type)
        values = collect_form_values(dict(form_data), helper)

        for field_name, items in state.current_nested_items.items():
            if items:
                values[field_name] = items

        errors = []
        if state.profile == "miappe":
            errors = validate_data(values, entity_type, version=facade.version)

        error_list = [{"field": e.field, "message": e.message, "rule": e.rule} for e in errors]

        return templates.TemplateResponse(
            request,
            "components/validation_result.html",
            {"valid": len(errors) == 0, "errors": error_list},
        )

    return app


# -------------------------------------------------------------------------
# Private helper functions for route handlers
# -------------------------------------------------------------------------


def _render_entity_form(
    request: Request,
    templates: Jinja2Templates,
    facade: ProfileFacade,
    helper: Any,
    entity_type: str,
    node_id: str,
    instance: Any,
    success_message: str,
    state: AppState | None = None,
) -> HTMLResponse:
    """Render entity form after successful create/update."""
    fields = get_field_data(helper)
    values = instance.model_dump(exclude_none=True) if hasattr(instance, "model_dump") else {}

    auto_fields = set()
    if "miappe_version" in helper.all_fields:
        values["miappe_version"] = facade.version
        auto_fields.add("miappe_version")

    # Build inline table data if state is available
    inline_tables = {}
    if state:
        inline_tables = build_inline_tables(state, facade, entity_type)

    response = templates.TemplateResponse(
        request,
        "partials/form.html",
        {
            "entity_type": entity_type,
            "is_edit": True,
            "node_id": node_id,
            "description": helper.description,
            "ontology_term": helper.ontology_term,
            "required_fields": [f for f in fields if f["required"]],
            "optional_fields": [f for f in fields if not f["required"] and not is_nested_field(f)],
            "nested_fields": [f for f in fields if is_nested_field(f)],
            "values": values,
            "auto_fields": auto_fields,
            "success_message": success_message,
            "inline_tables": inline_tables,
        },
    )
    response.headers["HX-Trigger"] = (
        "entityCreated" if "Created" in success_message else "entityUpdated"
    )
    return response


def _render_form_with_errors(
    request: Request,
    templates: Jinja2Templates,
    facade: ProfileFacade,
    helper: Any,
    entity_type: str,
    node_id: str | None,
    values: dict,
    error: ValidationError,
) -> HTMLResponse:
    """Render form with validation errors."""
    errors = format_validation_errors(error)
    fields = get_field_data(helper)

    auto_fields = set()
    if "miappe_version" in helper.all_fields:
        values["miappe_version"] = facade.version
        auto_fields.add("miappe_version")

    return templates.TemplateResponse(
        request,
        "partials/form.html",
        {
            "entity_type": entity_type,
            "is_edit": node_id is not None,
            "node_id": node_id,
            "description": helper.description,
            "ontology_term": helper.ontology_term,
            "required_fields": [f for f in fields if f["required"]],
            "optional_fields": [f for f in fields if not f["required"] and not is_nested_field(f)],
            "nested_fields": [f for f in fields if is_nested_field(f)],
            "values": values,
            "auto_fields": auto_fields,
            "error_message": f"Validation error: {errors}",
        },
    )


app = create_app()


def run_ui(host: str = "127.0.0.1", port: int = 8080) -> None:
    """Run the Metaseed web interface."""
    import uvicorn

    uvicorn.run(app, host=host, port=port)
